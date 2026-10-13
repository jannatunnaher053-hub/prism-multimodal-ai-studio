import base64
import io
import json
import os
import tempfile
from pathlib import Path

import requests


OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434")


def get_gemini_key():
    key = os.getenv("GEMINI_API_KEY", "").strip()

    if not key:
        try:
            import streamlit as st
            key = str(st.secrets.get("GEMINI_API_KEY", "")).strip()
        except Exception:
            pass

    return key


def gemini_generate(prompt, image_bytes=None):
    key = get_gemini_key()

    if not key:
        return (
            "Gemini API key পাওয়া যায়নি। "
            "Streamlit Cloud → App settings → Secrets-এ "
            "GEMINI_API_KEY যোগ করো।"
        )

    model = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
    parts = [{"text": prompt}]

    if image_bytes:
        from PIL import Image, ImageOps

        try:
            image = Image.open(io.BytesIO(image_bytes))
            image = ImageOps.exif_transpose(image)
            image.thumbnail((1024, 1024))

            if image.mode != "RGB":
                image = image.convert("RGB")

            buffer = io.BytesIO()
            image.save(buffer, format="JPEG", quality=80, optimize=True)
            compressed_bytes = buffer.getvalue()

            parts.append({
                "inline_data": {
                    "mime_type": "image/jpeg",
                    "data": base64.b64encode(compressed_bytes).decode("utf-8"),
                }
            })

        except Exception as exc:
            return f"Image processing failed: {exc}"


    url = (
        "https://generativelanguage.googleapis.com/"
        f"v1beta/models/{model}:generateContent"
    )

    try:
        response = requests.post(
            url,
            headers={
                "x-goog-api-key": key,
                "Content-Type": "application/json",
            },
            json={"contents": [{"parts": parts}]},
            timeout=120,
        )

        if not response.ok:
            return (
                f"Gemini API error ({response.status_code}): "
                f"{response.text[:1000]}"
            )

        data = response.json()
        candidates = data.get("candidates", [])

        if not candidates:
            return "Gemini কোনো উত্তর দেয়নি। অন্য একটি prompt দিয়ে চেষ্টা করো।"

        answer_parts = candidates[0].get("content", {}).get("parts", [])
        answer = "\n".join(
            part.get("text", "")
            for part in answer_parts
            if part.get("text")
        ).strip()

        return answer or "Gemini একটি খালি উত্তর দিয়েছে।"

    except requests.RequestException as exc:
        return f"Gemini API-তে সংযোগ করা যায়নি: {exc}"
    except Exception as exc:
        return f"উত্তর তৈরি করতে সমস্যা হয়েছে: {exc}"


def ollama_status():
    if get_gemini_key():
        return "Gemini API configured. Cloud AI is ready to test."

    try:
        r = requests.get(f"{OLLAMA_URL}/api/tags", timeout=2)
        r.raise_for_status()
        names = [m.get("name", "") for m in r.json().get("models", [])]
        return "Connected to Ollama. Models: " + (
            ", ".join(names) if names else "none pulled yet"
        )
    except Exception:
        return (
            "Gemini API key not configured and Ollama not detected. "
            "Configure GEMINI_API_KEY in Streamlit Cloud Secrets."
        )


def ollama_chat(model, prompt):
    if get_gemini_key():
        return gemini_generate(prompt)

    try:
        r = requests.post(
            f"{OLLAMA_URL}/api/generate",
            json={"model": model, "prompt": prompt, "stream": False},
            timeout=300,
        )
        r.raise_for_status()
        return r.json().get("response", "").strip() or (
            "The model returned an empty response."
        )
    except requests.exceptions.ConnectionError:
        return (
            "Cannot connect to Ollama. Configure GEMINI_API_KEY in "
            "Streamlit Cloud Secrets to use cloud AI."
        )
    except requests.exceptions.HTTPError as e:
        return f"Ollama request failed: {e}"


def image_question(model, prompt, image_bytes):
    if get_gemini_key():
        return gemini_generate(prompt, image_bytes=image_bytes)

    encoded = base64.b64encode(image_bytes).decode("utf-8")

    try:
        r = requests.post(
            f"{OLLAMA_URL}/api/generate",
            json={
                "model": model,
                "prompt": prompt,
                "images": [encoded],
                "stream": False,
            },
            timeout=300,
        )
        r.raise_for_status()
        return r.json().get("response", "").strip()
    except requests.exceptions.ConnectionError:
        return "Cannot connect to Ollama. Configure Gemini API for cloud use."
    except requests.exceptions.HTTPError as e:
        return f"Vision model request failed: {e}"


def extract_document_text(filename, data):
    ext = Path(filename).suffix.lower()
    if ext == ".txt":
        return data.decode("utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    raise ValueError("Supported document types: PDF, DOCX, TXT.")

def create_presentation(title, outline):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = "Created with Prism Multimodal AI Studio"
    sections = [s.strip() for s in outline.split("===") if s.strip()]
    for section in sections:
        lines = [line.strip() for line in section.splitlines() if line.strip()]
        if not lines:
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = lines[0].lstrip("#* ")
        tf = slide.placeholders[1].text_frame
        tf.clear()
        bullet_count = 0
        for line in lines[1:]:
            text = line.lstrip("-•* ").strip()
            if not text:
                continue
            p = tf.paragraphs[0] if bullet_count == 0 else tf.add_paragraph()
            p.text = text
            p.level = 0
            p.font.size = Pt(22)
            bullet_count += 1
        if bullet_count == 0:
            tf.text = "Add details for this slide."
    out = Path(tempfile.gettempdir()) / "prism_presentation.pptx"
    prs.save(out)
    return str(out)

def transcribe_audio_local(audio_bytes, filename, model_size="base"):
    from faster_whisper import WhisperModel
    suffix = Path(filename).suffix or ".wav"
    temp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    try:
        temp.write(audio_bytes)
        temp.close()
        # CPU mode works on ordinary Windows PCs; GPU/CUDA can be configured later.
        model = WhisperModel(model_size, device="cpu", compute_type="int8")
        segments, info = model.transcribe(temp.name, beam_size=5)
        return " ".join(seg.text.strip() for seg in segments).strip()
    finally:
        try:
            os.unlink(temp.name)
        except OSError:
            pass

def speak_text_local(text):
    import pyttsx3
    out = Path(tempfile.gettempdir()) / "prism_voice_response.wav"
    engine = pyttsx3.init()
    engine.save_to_file(text[:5000], str(out))
    engine.runAndWait()
    engine.stop()
    if not out.exists():
        raise RuntimeError("Speech file was not created.")
    return str(out)

def generate_image_local(prompt):
    import torch
    from diffusers import StableDiffusionPipeline
    model_id = os.getenv("PRISM_IMAGE_MODEL", "stable-diffusion-v1-5/stable-diffusion-v1-5")
    device = "cuda" if torch.cuda.is_available() else "cpu"
    dtype = torch.float16 if device == "cuda" else torch.float32
    pipe = StableDiffusionPipeline.from_pretrained(model_id, torch_dtype=dtype, safety_checker=None)
    pipe = pipe.to(device)
    if device == "cpu":
        pipe.enable_attention_slicing()
    image = pipe(prompt, num_inference_steps=20).images[0]
    out = Path(tempfile.gettempdir()) / "prism_generated.png"
    image.save(out)
    return str(out)

def create_storyboard_video(title, script):
    from PIL import Image, ImageDraw, ImageFont
    import imageio.v2 as imageio
    import numpy as np
    scenes = [s.strip() for s in script.split("\n\n") if s.strip()]
    if not scenes:
        raise ValueError("Enter at least one scene paragraph.")
    width, height = 1280, 720
    font_path = "C:/Windows/Fonts/arial.ttf"
    title_font = ImageFont.truetype(font_path, 48) if os.path.exists(font_path) else ImageFont.load_default()
    body_font = ImageFont.truetype(font_path, 36) if os.path.exists(font_path) else ImageFont.load_default()
    frames = []
    for idx, scene in enumerate(scenes):
        img = Image.new("RGB", (width, height), (18, 24, 38))
        draw = ImageDraw.Draw(img)
        draw.text((70, 60), title[:70], font=title_font, fill=(255, 255, 255))
        draw.rounded_rectangle((70, 160, width-70, height-90), radius=24, fill=(34, 45, 66), outline=(105, 130, 170), width=3)
        words = scene.split()
        lines, line = [], ""
        for word in words:
            candidate = (line + " " + word).strip()
            if len(candidate) > 42 and line:
                lines.append(line)
                line = word
            else:
                line = candidate
        if line:
            lines.append(line)
        y = 250
        for text in lines[:7]:
            draw.text((120, y), text, font=body_font, fill=(240, 244, 250))
            y += 58
        draw.text((width-180, height-65), f"{idx+1}/{len(scenes)}", font=body_font, fill=(170, 190, 220))
        arr = np.asarray(img)
        # Hold each scene for 2.5 seconds at 12 fps.
        frames.extend([arr] * 30)
    out = Path(tempfile.gettempdir()) / "prism_storyboard.mp4"
    imageio.mimwrite(out, frames, fps=12, codec="libx264", quality=7)
    return str(out)
